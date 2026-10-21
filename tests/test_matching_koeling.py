import io

import openpyxl
import pandas as pd
import pytest

from matching_koeling import (
    BudgetLink,
    IBItem,
    OfferteItem,
    HAS_RAPIDFUZZ,
    _nh,
    _norm_code,
    _normalize_jum_ko_code,
    _parse_buffet_manual_code,
    _parse_nrfp_rows,
    _parse_offerte_rows,
    _pdf_cell,
    _stringify_manual_code,
    apply_art_nr_mapping,
    build_budget_link,
    build_buffetten_link,
    CATEGORIE_GROEP_LABEL,
    NrfpItem,
    bidsheet_leverancier,
    build_installatie_df,
    build_offerte_vergelijking,
    guess_leverancier,
    match_koeling_installatie,
    parse_buiten_rfp,
    parse_ib,
    parse_offerte,
    vergelijking_samenvatting,
    vergelijking_to_xlsx,
)
from vergelijking import unique_namen, vergelijking_regels_df


def _xlsx_bytes(sheet_name: str, rows: list) -> bytes:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = sheet_name
    for row in rows:
        ws.append(row)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


# ─── normalization helpers ──────────────────────────────────────────────────

def test_nh_normalizes_header_labels():
    assert _nh("Omschrijving Item") == "omschrijvingitem"
    assert _nh(None) == ""
    assert _nh("  Art.Nr.Jumbo  ") == "artnrjumbo"


@pytest.mark.parametrize("value,expected", [
    (None, None),
    ("Jum-ko-0470", "Jum-ko-0470"),
    (470.0, "470"),
    (19.81, "19.81"),
])
def test_stringify_manual_code(value, expected):
    assert _stringify_manual_code(value) == expected


@pytest.mark.parametrize("raw,expected", [
    ("Jum-ko-0470", "Jum-ko-0470"),
    ("jum-ko-470", "Jum-ko-0470"),
    ("470", "Jum-ko-0470"),
    ("19.81", None),
])
def test_normalize_jum_ko_code(raw, expected):
    assert _normalize_jum_ko_code(raw) == expected


@pytest.mark.parametrize("raw,expected", [
    ("19.8", "19.80"),
    (19.81, "19.81"),
    ("not-a-number", None),
])
def test_norm_code(raw, expected):
    assert _norm_code(raw) == expected


def test_parse_buffet_manual_code():
    assert _parse_buffet_manual_code(None) is None
    assert _parse_buffet_manual_code(19.81) == ("19.81", None)
    assert _parse_buffet_manual_code("19.80-4drs") == ("19.80", "4")
    assert _parse_buffet_manual_code("Jum-ko-0470") is None


# ─── Excel parsing ──────────────────────────────────────────────────────────

def test_parse_offerte_reads_items_and_manual_code():
    rows = [
        ["Categorie koeling"],
        ["Subgroep", "Omschrijving Item", "Eenheid", "Prijs materiaal per eenheid",
         "Arbeidskosten per eenheid", "Aantal", "Art.Nr.Jumbo"],
        ["Sub A", "Gas Cooler 133kW", "stuks", 1000.0, 50.0, 2, 470],
    ]
    items = parse_offerte(_xlsx_bytes("Offerte", rows), "Offerte")

    assert len(items) == 1
    item = items[0]
    assert item.categorie == "Categorie koeling"
    assert item.omschrijving == "Gas Cooler 133kW"
    assert item.aantal == 2
    assert item.art_manual == "470"  # numeric cell stringified


def test_parse_offerte_reads_aantal_in_normwinkel_header():
    rows = [
        ["Categorie koelcellen", "Omschrijving Item", "Eenheid", "Prijs materiaal per eenheid",
         "Arbeidskosten per eenheid", "Aantal in normwinkel"],
        [None, "Verdamper koelcel 3000 W CO₂", "stuks", 900.0, 56.0, 2],
    ]
    items = parse_offerte(_xlsx_bytes("Bidsheet K", rows), "Bidsheet K")

    assert len(items) == 1
    assert items[0].categorie == "Categorie koelcellen"
    assert items[0].aantal == 2


def test_parse_offerte_falls_back_to_netto_netto_prijs():
    rows = [
        ["Categorie verdamper vleesafdeling", "Omschrijving Item", "Eenheid", "Prijs materiaal per eenheid",
         "Arbeidskosten per eenheid", "Netto netto prijs per eenheid (arbeid & materiaal)", "Aantal in normwinkel"],
        [None, "Ruimteverdamper vleesafdeling CO₂", "st", None, None, 4107, 1],
        [None, "Verdamper met eigen prijs", "st", 900.0, None, 950.0, 1],
    ]
    items = parse_offerte(_xlsx_bytes("Bidsheet K", rows), "Bidsheet K")

    assert items[0].prijs_materiaal == 4107
    assert items[1].prijs_materiaal == 900.0  # materiaal filled in -> netto not used


def test_buiten_rfp_lines_become_nrfp_items_not_offerte_items():
    rows = [
        ["Categorie demontage", "Omschrijving Item", "Eenheid", "Prijs materiaal per eenheid",
         "Arbeidskosten per eenheid", "Netto netto prijs per eenheid (arbeid & materiaal)", "Aantal in normwinkel"],
        ["Afpompen", "Koudemiddel afpompen", "kg", 12, None, 12, 10],
        ["Buiten RFP", "Omschrijving Item"],  # short header: no Aantal/Prijs labels
        ["Buiten RFP", "Afblazen bestaande installatie", "st", 468, None, 468, 1],
        ["Buiten RFP", None],
    ]
    data = _xlsx_bytes("Bidsheet K", rows)

    assert [i.omschrijving for i in parse_offerte(data, "Bidsheet K")] == ["Koudemiddel afpompen"]
    buiten = parse_buiten_rfp(data, "Bidsheet K")
    assert len(buiten) == 1
    assert buiten[0].omschrijving == "Afblazen bestaande installatie"
    assert buiten[0].totaal == 468.0


def test_parse_ib_skips_nrfp_rows():
    rows = [
        ["Art.Nr.Jumbo", "Manual nr.", "Art.Nr.Leverancier", "Groep", "Omschrijving Item", "Eenheid", "Aantal",
         "Materiaalkosten per eenheid"],
        ["Jum-ko-0470", None, "SUP-1", "Koeling", "Gas Cooler 133kW", "stuks", 2, 1000.0],
        ["NRFP-1", None, None, "Koeling", "Niet relevant", "stuks", 1, 10.0],
    ]
    items = parse_ib(_xlsx_bytes("IB", rows), "IB")

    assert len(items) == 1
    assert items[0].art_nr_jumbo == "Jum-ko-0470"
    assert items[0].materiaalkosten == 1000.0


# ─── matching ────────────────────────────────────────────────────────────────

def _offerte_item(**overrides) -> OfferteItem:
    base = dict(
        row=1, categorie="Koeling", subgroep="Sub A", omschrijving="Gas Cooler 133kW",
        eenheid="stuks", prijs_materiaal=1000.0, arbeid=50.0, aantal=2, art_manual=None,
    )
    base.update(overrides)
    return OfferteItem(**base)


def _ib_item(**overrides) -> IBItem:
    base = dict(
        row=5, art_nr_jumbo="Jum-ko-0470", manual_nr=None, art_nr_leverancier=None,
        groep="Koeling", omschrijving="Gas Cooler 133kW", eenheid="stuks", aantal=2,
        materiaalkosten=1000.0,
    )
    base.update(overrides)
    return IBItem(**base)


def test_match_koeling_installatie_manual_exact():
    offerte = [_offerte_item(art_manual="Jum-ko-0470")]
    ib = [_ib_item()]

    matches, used_ib = match_koeling_installatie(offerte, ib)

    assert matches[0] == (0, 100.0, "handmatig")
    assert used_ib == {0}


def test_match_koeling_installatie_normalizes_bare_number():
    offerte = [_offerte_item(art_manual="470")]
    ib = [_ib_item()]

    matches, _ = match_koeling_installatie(offerte, ib)

    assert matches[0][0] == 0
    assert matches[0][2] == "handmatig"


def test_match_koeling_installatie_no_match_without_rapidfuzz_or_manual_code():
    offerte = [_offerte_item(art_manual=None, omschrijving="Totally unrelated widget")]
    ib = [_ib_item()]

    matches, used_ib = match_koeling_installatie(offerte, ib, fuzzy_threshold=999)

    assert matches == {}
    assert used_ib == set()


@pytest.mark.skipif(not HAS_RAPIDFUZZ, reason="rapidfuzz not installed")
def test_match_koeling_installatie_fuzzy_fallback():
    offerte = [_offerte_item(art_manual=None)]
    ib = [_ib_item()]

    matches, used_ib = match_koeling_installatie(offerte, ib, fuzzy_threshold=40)

    assert matches[0][0] == 0
    assert matches[0][2] == "fuzzy"
    assert used_ib == {0}


def test_apply_art_nr_mapping_does_not_override_existing():
    offerte = [
        _offerte_item(row=1, omschrijving="Gas Cooler 133kW", art_manual=None),
        _offerte_item(row=2, omschrijving="Already Set", art_manual="Jum-ko-0001"),
    ]
    mapping = {"gas cooler 133kw": "Jum-ko-0470", "already set": "Jum-ko-9999"}

    filled = apply_art_nr_mapping(offerte, mapping)

    assert filled == 1
    assert offerte[0].art_manual == "Jum-ko-0470"
    assert offerte[1].art_manual == "Jum-ko-0001"  # untouched — offerte value wins


# ─── df building / status ───────────────────────────────────────────────────

def test_build_installatie_df_flags_price_and_quantity_differences():
    offerte = [
        _offerte_item(row=1, prijs_materiaal=1000.0, aantal=2),   # price mismatch
        _offerte_item(row=2, prijs_materiaal=1000.0, aantal=3),   # qty mismatch
        _offerte_item(row=3, prijs_materiaal=1000.0, aantal=2),   # ok
        _offerte_item(row=4, art_manual=None),                    # unmatched
    ]
    ib = [
        _ib_item(row=5, materiaalkosten=900.0, aantal=2),
        _ib_item(row=6, materiaalkosten=1000.0, aantal=2),
        _ib_item(row=7, materiaalkosten=1000.0, aantal=2),
    ]
    matches = {0: (0, 100.0, "handmatig"), 1: (1, 100.0, "handmatig"), 2: (2, 100.0, "handmatig")}
    # Aantal is only compared via the Budget-sheet link (build_installatie_df's
    # own aantal_i always comes from budget_link/buffetten_link, never straight
    # from ib_item.aantal) — so row 1 needs a link with a differing aantal.
    budget_link = {6: BudgetLink(budget_row=999, code=None, omschrijving=None, aantal=2)}

    df = build_installatie_df(offerte, ib, matches, budget_link=budget_link)

    assert list(df["_status"]) == ["prijs_afwijking", "aantal_afwijking", "ok", "unmatched"]


def test_build_installatie_df_uses_budget_link_omschrijving_and_aantal():
    offerte = [_offerte_item(row=1, prijs_materiaal=1000.0, aantal=2)]
    ib = [_ib_item(row=5, materiaalkosten=1000.0, aantal=2)]
    matches = {0: (0, 100.0, "handmatig")}
    budget_link = {5: BudgetLink(budget_row=1513, code="0470", omschrijving="Budget omschrijving", aantal=2)}

    df = build_installatie_df(offerte, ib, matches, budget_link=budget_link)

    assert df.loc[0, "IB omschrijving"] == "Budget omschrijving"
    assert df.loc[0, "_status"] == "ok"


# ─── budget linking ──────────────────────────────────────────────────────────

def _budget_df(rows: dict) -> pd.DataFrame:
    """rows: {1-based row number: (code, desc, aantal)} -> an 11-column frame
    shaped like the real Budget sheet (code=G/6, desc=H/7, aantal=K/10)."""
    max_row = max(rows)
    data = [[None] * 11 for _ in range(max_row)]
    for r, (code, desc, aantal) in rows.items():
        data[r - 1][6] = code
        data[r - 1][7] = desc
        data[r - 1][10] = aantal
    return pd.DataFrame(data)


def test_build_budget_link_joins_on_numeric_suffix():
    ib = [_ib_item(row=5, art_nr_jumbo="Jum-ko-0470")]
    bdf = _budget_df({1: ("COOL-001", "0470 Gas Cooler Ec fans 133kW", 3.5)})

    links = build_budget_link(bdf, ib, row_start=1, row_end=1)

    assert links[5] == BudgetLink(budget_row=1, code="COOL-001", omschrijving="0470 Gas Cooler Ec fans 133kW", aantal=3.5)


def test_build_budget_link_none_bdf_returns_empty():
    assert build_budget_link(None, [_ib_item()]) == {}


@pytest.mark.skipif(not HAS_RAPIDFUZZ, reason="rapidfuzz not installed")
def test_build_budget_link_falls_back_when_matched_row_has_no_aantal():
    # Suffix-joined row (0470) has no aantal — IB actually selected a
    # differently-numbered but similarly-described variant on another row.
    ib = [_ib_item(row=5, art_nr_jumbo="Jum-ko-0470", omschrijving="Bypass valve 5")]
    bdf = _budget_df({
        1: ("COOL-001", "0470 Bypass valve 5", 0),
        2: ("COOL-002", "Bypass valve 4", 3.0),
    })

    links = build_budget_link(bdf, ib, row_start=1, row_end=2)

    assert links[5].budget_row == 2
    assert links[5].aantal == 3.0


def test_build_budget_link_keeps_zero_aantal_when_no_better_match():
    ib = [_ib_item(row=5, art_nr_jumbo="Jum-ko-0470", omschrijving="Bypass valve 5")]
    bdf = _budget_df({1: ("COOL-001", "0470 Bypass valve 5", 0)})

    links = build_budget_link(bdf, ib, row_start=1, row_end=1)

    assert links[5].budget_row == 1
    assert not links[5].aantal


def test_build_buffetten_link_manual_code_takes_priority():
    offerte = [_offerte_item(row=1, categorie="Categorie buffetten", art_manual="19.81")]
    bdf = _budget_df({1442: ("19.81", "19.81 Saladiere buffet", 1.0)})

    links = build_buffetten_link(bdf, offerte)

    link, source = links[1]
    assert source == "handmatig"
    assert link.budget_row == 1442


def test_build_buffetten_link_falls_back_to_deurs_heuristic():
    offerte = [_offerte_item(row=1, categorie="Categorie buffetten", art_manual=None,
                              omschrijving="Koelbuffet 2-deurs")]
    bdf = _budget_df({1442: (None, "Prijs variant 2 drs koelbuffet", 1.0)})

    links = build_buffetten_link(bdf, offerte)

    link, source = links[1]
    assert source == "budget-tekst"
    assert link.budget_row == 1442


# ─── offertevergelijking ────────────────────────────────────────────────────

def _match_df(rows: list) -> pd.DataFrame:
    """Minimal build_installatie_df-shaped frame: one dict per offerte item."""
    base = {"Art.Nr.Jumbo": None, "IB omschrijving": None, "IB Aantal": None, "IB Prijs p.e.": None,
            "Match": "geen match", "Match %": ""}
    return pd.DataFrame([{**base, **r} for r in rows])


def _vergelijking():
    frimex_items = [
        _offerte_item(row=10, categorie="Categorie koelcellen", omschrijving="Verdamper koelcel 3000 W",
                      prijs_materiaal=900.0, arbeid=100.0, aantal=2),
        _offerte_item(row=11, categorie="Categorie vriescellen", omschrijving="Verdamper vriescel 600 W",
                      prijs_materiaal=500.0, arbeid=0.0, aantal=1),
        _offerte_item(row=20, categorie="Categorie buffetten", omschrijving="Koelbuffet 4-deurs",
                      prijs_materiaal=6000.0, arbeid=None, aantal=1),
        _offerte_item(row=21, categorie="Categorie buffetten", omschrijving="Niet gekozen variant",
                      prijs_materiaal=7000.0, arbeid=None, aantal=None),
    ]
    frimex_df = _match_df([
        {"Art.Nr.Jumbo": "Jum-ko-0030", "IB Aantal": 2, "IB Prijs p.e.": 950.0, "Match": "handmatig"},
        {"Art.Nr.Jumbo": "Jum-ko-0080", "IB Aantal": 1, "IB Prijs p.e.": 600.0, "Match": "handmatig"},
        {},
        {},
    ])
    carrier_items = [
        # different wording, but matched to the same IB article
        _offerte_item(row=23, categorie="Categorie koelcellen", omschrijving="Verdamper koelcel 3000 W CO2",
                      prijs_materiaal=956.0, arbeid=None, aantal=2),
        # no IB match, but same template omschrijving as Frimex's matched line
        _offerte_item(row=31, categorie="Categorie vriescellen", omschrijving="Verdamper vriescel 600 W",
                      prijs_materiaal=450.0, arbeid=None, aantal=1),
    ]
    carrier_df = _match_df([
        {"Art.Nr.Jumbo": "Jum-ko-0030", "IB Aantal": 2, "IB Prijs p.e.": 950.0, "Match": "fuzzy", "Match %": "87%"},
        {},
    ])
    return build_offerte_vergelijking(
        [("Frimex", "frimex.xlsx", frimex_items, frimex_df, [NrfpItem("NRFP 1", "Huur container", 100.0, 2, 200.0)]),
         ("Carrier", "carrier.xlsx", carrier_items, carrier_df, [])],
        categorie_totalen={"Categorie buffetten": 5000.0},
        categorie_mapping={"Categorie buffetten": [("C.01.07", "Koelbuffetten")]},
    )


def test_vergelijking_aligns_offertes_per_ib_article():
    v = _vergelijking()

    # cellen grouped into one categorie; unselected variant dropped
    assert v.categorieen == [CATEGORIE_GROEP_LABEL, "Categorie buffetten"]
    assert len(v.regels) == 3
    koelcel, vriescel, buffet = v.regels
    assert koelcel.art_nr_jumbo == "Jum-ko-0030"
    assert koelcel.offertes == [(2, 1000.0), (2, 956.0)]
    assert "Carrier: fuzzy match 87%" in koelcel.opmerking
    # Carrier's unmatched line folded into Frimex's IB-article line by omschrijving
    assert vriescel.art_nr_jumbo == "Jum-ko-0080"
    assert vriescel.offertes == [(1, 500.0), (1, 450.0)]
    assert buffet.offertes == [(1, 6000.0), (None, None)]
    assert v.ib_categorie_totalen == {"Categorie buffetten": 5000.0}


def test_vergelijking_samenvatting_totals_and_laagste():
    sam = vergelijking_samenvatting(_vergelijking()).set_index("Categorie")

    cellen = sam.loc[CATEGORIE_GROEP_LABEL]
    assert cellen["IB (Budget)"] == pytest.approx(2 * 950 + 600)  # no Budget code -> sum of IB lines
    assert cellen["Frimex"] == pytest.approx(2500.0)
    assert cellen["Carrier"] == pytest.approx(2362.0)
    assert cellen["Laagste"] == "Carrier"
    assert sam.loc["Categorie buffetten", "IB (Budget)"] == 5000.0  # Budget subtotal wins
    assert sam.loc["Categorie buffetten", "Laagste"] == "Alleen Frimex"
    assert sam.loc["NRFP / Buiten RFP", "Laagste"] == ""
    assert sam.loc["Totaal incl. NRFP", "Frimex"] == pytest.approx(2500 + 6000 + 200)


def test_vergelijking_to_xlsx_writes_formula_workbook():
    wb = openpyxl.load_workbook(io.BytesIO(vergelijking_to_xlsx(_vergelijking())))

    assert wb.sheetnames == ["Samenvatting", "Vergelijking", "NRFP"]
    ws = wb["Vergelijking"]
    assert ws["A6"].value == CATEGORIE_GROEP_LABEL
    assert ws["I6"].value == 2 and ws["J6"].value == 1000.0           # Frimex aantal / prijs
    assert ws["K6"].value == '=IF(OR(I6="",J6=""),0,I6*J6)'          # Frimex totaal is a formula
    assert ws["Q6"].value.startswith("=IF((K6>0)+(O6>0)=0")          # Laagste formula
    s = wb["Samenvatting"]
    sam_rows = {s.cell(r, 1).value: r for r in range(1, s.max_row + 1)}
    buffet_row = sam_rows["Categorie buffetten"]
    assert s.cell(buffet_row, 2).value == 5000.0                     # Budget subtotal as value
    assert s.cell(buffet_row, 3).value.startswith("=SUMIFS(Vergelijking!$K$")


def test_vergelijking_to_xlsx_single_offerte_has_no_laagste():
    v = _vergelijking()
    v.namen, v.bestanden, v.nrfp = v.namen[:1], v.bestanden[:1], v.nrfp[:1]
    for r in v.regels:
        r.offertes = r.offertes[:1]
    wb = openpyxl.load_workbook(io.BytesIO(vergelijking_to_xlsx(v)))

    headers = {c.value for ws in wb for row in ws.iter_rows(max_row=12) for c in row if isinstance(c.value, str)}
    assert "Laagste" not in headers and "Laagste offerte" not in headers
    assert wb["Vergelijking"]["M4"].value == "Opmerking"  # right after the single offerte's 4 columns

    sam = vergelijking_samenvatting(v)
    assert "Laagste" not in sam.columns
    assert sam.set_index("Categorie").loc["Totaal incl. NRFP", "Frimex"] == pytest.approx(2500 + 6000 + 200)


def test_vergelijking_regels_df_mirrors_excel_layout():
    rdf = vergelijking_regels_df(_vergelijking())

    assert list(rdf.columns) == [
        "Categorie", "Art.Nr.Jumbo", "Omschrijving (offerte)", "IB omschrijving (Budget)", "Eenheid",
        "IB Aantal", "IB Prijs p.e.", "IB Totaal",
        "Frimex Aantal", "Frimex Prijs p.e.", "Frimex Totaal", "Frimex − IB",
        "Carrier Aantal", "Carrier Prijs p.e.", "Carrier Totaal", "Carrier − IB",
        "Laagste offerte", "Opmerking",
    ]
    koelcel = rdf.iloc[0]
    assert koelcel["Frimex Totaal"] == 2000.0 and koelcel["Frimex − IB"] == pytest.approx(100.0)
    assert koelcel["Laagste offerte"] == "Carrier"
    tot = rdf.iloc[-1]
    assert tot["Categorie"] == "Totaal"
    assert tot["Frimex Totaal"] == pytest.approx(2000 + 500 + 6000)
    assert tot["IB Totaal"] == pytest.approx(2 * 950 + 600)


def test_unique_namen_suffixes_duplicates():
    assert unique_namen(["Frimex", "Carrier", "Frimex", "Frimex"]) == ["Frimex", "Carrier", "Frimex (2)", "Frimex (3)"]


def test_bidsheet_leverancier_reads_leverancier_line():
    rows = [
        ["Algemene informatie:"],
        ["Leverancier", "Carrier Commercial Refrigeration"],
        ["Categorie koelcellen", "Omschrijving Item", "Aantal"],
    ]
    assert bidsheet_leverancier(_xlsx_bytes("Bidsheet K", rows), "bidsheet.xlsx") == "Carrier"
    rows[1] = ["Leverancier", "Koeltechniek Jansen B.V."]
    assert bidsheet_leverancier(_xlsx_bytes("Bidsheet K", rows), "bidsheet.xlsx") == "Koeltechniek Jansen B.V."
    # no Leverancier line -> filename
    assert bidsheet_leverancier(_xlsx_bytes("Bidsheet K", rows[2:]), "Frimex bidsheet.xlsx") == "Frimex"


def test_guess_leverancier_from_filename():
    assert guess_leverancier("21011533 0.1 Jumbo Eibergen, Frimex KT bidsheet.xlsx") == "Frimex"
    assert guess_leverancier("Installatie bidsheet KT HVAC en W - CCR  Jumbo Eibergen.xlsx") == "Carrier"
    assert guess_leverancier("Offerte Koeltechniek Jansen.xlsx") == "Offerte Koeltechniek Jansen"


# ─── PDF bidsheets ──────────────────────────────────────────────────────────

@pytest.mark.parametrize("raw, expected", [
    ("€ 1 2.463,15", 12463.15),   # stray spaces inside the digits
    ("€ 7 ,31", 7.31),
    ("1 20", 120.0),              # aantal
    ("2", 2.0),
    ("€ -", 0.0),
    ("-", 0.0),
    ("", None),
    (None, None),
    ("Verdamper koelcel 1000 W CO₂", "Verdamper koelcel 1000 W CO₂"),
    ("st", "st"),
])
def test_pdf_cell_converts_dutch_amounts(raw, expected):
    assert _pdf_cell(raw) == expected


def test_pdf_rows_parse_like_excel_and_stop_at_nrfp_page():
    raw = [
        ["Leverancier", "Frimex Installatie B.V."],
        ["Categorie koelcellen", "Omschrijving Item", "Eenheid", "Prijs materiaal\nper eenheid",
         "Arbeidskosten\nper eenheid", "Netto netto prijs\nper eenheid\n(materiaal en\narbeid)", "Aantal",
         "Totaalprijs\nwinkel in €", "Totaalprijs\ncategorie in €", "Opmerkingen"],
        ["Verdamper koelcel", "Verdamper koelcel 3000 W CO₂", "st", "€ 1 .038,01", "€ 149,54", "€ 1 .187,55", "1",
         "€ 1 .187,55", "", ""],
        ["Celpanelen koelcel", "Celpanelen koelcel hoogte 2400mm", "m2", "€ 141,56", "€ -", "€ 141,56", "1 03",
         "€ 1 4.580,37", "", ""],
        # next page: the table continues without its own header row
        ["Verdamper koelcel", "Verdamper koelcel 5000 W CO₂", "st", "€ 1 .269,97", "€ 149,54", "€ 1 .419,51", "",
         "€ -", "", ""],
        ["Positions not included in RFP", None, None, None, None],
        ["", "Apparatuur/meubelen", "Prijs", "Aantal", "Totaal"],
        ["NRFP 0010", "", "", "", "€ -"],
        ["NRFP 0170", "Drukopnemers meerprijs", "€ 2.872,16", "1", "€ 2.872,16"],
    ]
    rows = [(i, [_pdf_cell(v) for v in r]) for i, r in enumerate(raw, 1)]

    items = _parse_offerte_rows(rows)
    assert [i.omschrijving for i in items] == [
        "Verdamper koelcel 3000 W CO₂", "Celpanelen koelcel hoogte 2400mm", "Verdamper koelcel 5000 W CO₂"]
    assert items[0].prijs_materiaal == pytest.approx(1038.01)
    assert items[0].arbeid == pytest.approx(149.54)
    assert items[1].aantal == 103.0
    assert items[2].aantal is None

    nrfp = _parse_nrfp_rows(rows)
    assert [(n.code, n.totaal) for n in nrfp] == [("NRFP 0170", 2872.16)]
